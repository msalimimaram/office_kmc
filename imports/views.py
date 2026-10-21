import secrets
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.shortcuts import render, redirect, get_object_or_404
from openpyxl import load_workbook

from claims.models import Claim
from organizations.models import Person
from .forms import ClaimImportForm
from .models import ImportBatch
from auditlog.models import AuditEvent

# Until insurer accounts are linked to a specific Insurer record, they must not
# be allowed to import claims across insurer boundaries.
IMPORT_ROLES = {"system_admin", "department_head", "insurance_specialist"}


def _can_import(user):
    return user.is_superuser or user.role in IMPORT_ROLES


def _parse_amount(value, line, key):
    if value in (None, ""):
        return Decimal("0")
    try:
        amount = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f"ردیف {line}: ستون {key} باید عددی باشد.")
    if not amount.is_finite() or amount < 0:
        raise ValueError(f"ردیف {line}: ستون {key} باید عددی معتبر و نامنفی باشد.")
    return amount


def _parse_date(value, line, key):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip()[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        raise ValueError(f"ردیف {line}: تاریخ ستون {key} معتبر نیست؛ قالب YYYY-MM-DD استفاده شود.")


@login_required
def claim_import(request):
    if not _can_import(request.user):
        raise PermissionDenied
    if request.method == "POST":
        form = ClaimImportForm(request.POST, request.FILES)
        if form.is_valid():
            upload = form.cleaned_data["file"]
            wb = None
            try:
                wb = load_workbook(upload, read_only=True, data_only=True)
                ws = wb.active
                rows = ws.iter_rows(values_only=True)
                headers = next(rows, None)
                if not headers:
                    form.add_error("file", "فایل اکسل خالی است.")
                    return render(request, "imports/claim_import.html", {"form": form})
                headers = [str(v).strip().lower() if v is not None else "" for v in headers]
                required = {"claim_number", "national_id", "service_type"}
                if not required.issubset(set(headers)):
                    form.add_error("file", "ستون‌های اجباری: claim_number، national_id و service_type هستند.")
                    return render(request, "imports/claim_import.html", {"form": form})
                idx = {name: headers.index(name) for name in headers if name}
                data, errors, seen = [], [], set()
                for line, row in enumerate(rows, start=2):
                    if not any(v is not None for v in row):
                        continue

                    def val(key):
                        i = idx.get(key)
                        return row[i] if i is not None and i < len(row) else None

                    claim_no = str(val("claim_number") or "").strip()
                    nid = str(val("national_id") or "").strip()
                    if nid.endswith(".0"):
                        nid = nid[:-2]
                    service = str(val("service_type") or "").strip()
                    if not claim_no or not nid or not service:
                        errors.append(f"ردیف {line}: شناسه خسارت، کد ملی و نوع خدمت الزامی است.")
                        continue
                    if len(claim_no) > 100:
                        errors.append(f"ردیف {line}: شناسه خسارت بیش از حد مجاز طولانی است.")
                        continue
                    if claim_no in seen or Claim.objects.filter(claim_number=claim_no).exists():
                        errors.append(f"ردیف {line}: شناسه خسارت {claim_no} تکراری است.")
                        continue
                    people = Person.objects.filter(national_id=nid, is_active=True)
                    # Duplicate national IDs across subsidiaries must never be
                    # resolved by an arbitrary .first() match.
                    person_count = people.count()
                    if person_count == 0:
                        errors.append(f"ردیف {line}: شخصی با کد ملی واردشده یافت نشد.")
                        continue
                    if person_count > 1:
                        errors.append(f"ردیف {line}: کد ملی در چند شرکت یافت شد؛ داده‌ها باید رفع ابهام شوند.")
                        continue
                    person = people.get()
                    try:
                        amounts = {
                            key: _parse_amount(val(key), line, key)
                            for key in ("requested_amount", "assessed_amount", "payable_amount", "paid_amount")
                        }
                        dates = {
                            key: _parse_date(val(key), line, key)
                            for key in ("service_date", "registered_date", "payment_date")
                        }
                    except ValueError as exc:
                        errors.append(str(exc))
                        continue
                    if amounts["paid_amount"] > amounts["payable_amount"]:
                        errors.append(f"ردیف {line}: مبلغ پرداختی نمی‌تواند از مبلغ قابل پرداخت بیشتر باشد.")
                        continue
                    if dates["payment_date"] and not amounts["paid_amount"]:
                        errors.append(f"ردیف {line}: تاریخ پرداخت بدون مبلغ پرداختی ثبت شده است.")
                        continue
                    seen.add(claim_no)
                    data.append({
                        "claim_number": claim_no,
                        "person_id": person.pk,
                        "service_type": service[:120],
                        **{key: str(value) for key, value in amounts.items()},
                        **{key: value.isoformat() if value else None for key, value in dates.items()},
                        "source_reference": upload.name[:200],
                    })
                batch = ImportBatch.objects.create(
                    batch_number="IMP-" + secrets.token_hex(5).upper(),
                    file=upload, uploaded_by=request.user, total_rows=len(data) + len(errors),
                    valid_rows=len(data), invalid_rows=len(errors),
                    status=ImportBatch.Status.VALIDATED,
                    result_report="\n".join(errors[:300]),
                    preview_data=data,
                )
                return render(request, "imports/claim_import.html", {
                    "form": ClaimImportForm(), "batch": batch, "preview_count": len(data),
                    "errors": errors[:100], "confirm_available": bool(data),
                })
            except Exception:
                # Keep implementation details and file contents out of user-visible errors.
                form.add_error("file", "فایل قابل پردازش نیست یا ساختار آن معتبر نیست. قالب فایل و ستون‌ها را بررسی کنید.")
            finally:
                if wb is not None:
                    wb.close()
    else:
        form = ClaimImportForm()
    return render(request, "imports/claim_import.html", {"form": form})


@login_required
def claim_import_confirm(request, batch_id):
    if request.method != "POST" or not _can_import(request.user):
        raise PermissionDenied
    batch = get_object_or_404(
        ImportBatch, pk=batch_id, uploaded_by=request.user, status=ImportBatch.Status.VALIDATED
    )
    with transaction.atomic():
        for row in batch.preview_data:
            item = dict(row)
            for field in ("service_date", "registered_date", "payment_date"):
                if item.get(field):
                    item[field] = datetime.fromisoformat(item[field]).date()
            Claim.objects.create(**item)
        batch.status = ImportBatch.Status.PROCESSED
        batch.preview_data = []
        batch.save(update_fields=["status", "preview_data"])
    AuditEvent.objects.create(
        actor=request.user, action="claim_excel_import", object_type="ImportBatch",
        object_id=str(batch.pk), summary=f"{batch.valid_rows} ردیف ثبت شد؛ {batch.invalid_rows} خطا"
    )
    messages.success(request, f"ثبت نهایی انجام شد: {batch.valid_rows} خسارت؛ {batch.invalid_rows} ردیف خطا.")
    return redirect("claim_import")
